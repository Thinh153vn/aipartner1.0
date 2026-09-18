from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
from pathlib import Path

OUT = Path("GROWTH_PARTNER_Hackathon_Presentation.pptx")
MASCOT = Path("src/main/resources/static/assets/img/mascot.png")

W, H = 13.333, 7.5
NAVY = RGBColor(17, 36, 61)
NAVY2 = RGBColor(28, 52, 76)
CYAN = RGBColor(22, 183, 201)
MINT = RGBColor(32, 199, 164)
CORAL = RGBColor(240, 106, 98)
GOLD = RGBColor(242, 184, 75)
INK = RGBColor(21, 38, 59)
MUTED = RGBColor(105, 125, 140)
PAPER = RGBColor(247, 251, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(215, 231, 236)
VIOLET = RGBColor(116, 104, 232)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)

def shape(slide, kind, x, y, w, h, fill=None, line=None, radius=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    return s

def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box

def richbox(slide, runs, x, y, w, h, size=18, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.05); tf.margin_right = Inches(.05)
    p = tf.paragraphs[0]; p.alignment = align
    for txt, col, bold in runs:
        r = p.add_run(); r.text = txt; r.font.name = "Aptos"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = col
    return box

def bg(slide, color=PAPER):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, .08, CYAN)

def footer(slide, n, dark=False):
    col = RGBColor(165, 190, 200) if dark else MUTED
    textbox(slide, "SYSTEM EXE  /  INTERNAL HACKATHON", .55, 7.12, 5, .18, 8, col, True, "Aptos", margin=0)
    textbox(slide, f"{n:02d}", 12.2, 7.08, .55, .2, 9, col, True, "Aptos", PP_ALIGN.RIGHT, margin=0)

def title(slide, kicker, heading, sub=None, dark=False):
    col = WHITE if dark else INK; muted = RGBColor(183, 208, 214) if dark else MUTED
    textbox(slide, kicker.upper(), .7, .55, 5.5, .22, 9, CYAN if dark else CYAN, True, "Aptos", margin=0)
    textbox(slide, heading, .7, .93, 11.7, .62, 30, col, True, "Aptos Display", margin=0)
    if sub: textbox(slide, sub, .72, 1.67, 10.5, .42, 12, muted, False, "Aptos", margin=0)

def card(slide, x, y, w, h, heading, body, accent=CYAN, icon=None, dark=False):
    fill = NAVY2 if dark else WHITE; border = RGBColor(55, 82, 105) if dark else LINE
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, border)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, .055, h, accent)
    if icon: textbox(slide, icon, x+.22, y+.24, .42, .35, 18, accent, True, "Aptos", margin=0)
    textbox(slide, heading, x+.72 if icon else x+.22, y+.22, w-.95 if icon else w-.44, .3, 14, WHITE if dark else INK, True, margin=0)
    textbox(slide, body, x+.22, y+.7, w-.44, h-.88, 10.5, RGBColor(192, 211, 218) if dark else MUTED, False, margin=0)

def add_note(slide, text):
    # Keep a short presenter cue in the notes area when supported by python-pptx.
    try:
        slide.notes_slide.notes_text_frame.text = text
    except Exception:
        pass

# 1. Cover
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, NAVY)
shape(s, MSO_SHAPE.OVAL, 8.8, -1.4, 6.2, 6.2, NAVY2)
shape(s, MSO_SHAPE.OVAL, 10.0, -.2, 3.8, 3.8, None, CYAN)
shape(s, MSO_SHAPE.OVAL, 10.7, .5, 2.4, 2.4, None, MINT)
if MASCOT.exists(): s.shapes.add_picture(str(MASCOT), Inches(10.15), Inches(3.3), width=Inches(1.65), height=Inches(1.65))
textbox(s, "SYSTEM EXE  /  INTERNAL HACKATHON", .75, .72, 5.5, .2, 10, CYAN, True, margin=0)
textbox(s, "GROWTH\nPARTNER", .72, 1.55, 6.4, 1.55, 42, WHITE, True, "Aptos Display", margin=0)
textbox(s, "AIワークスペース for BrSE", .78, 3.35, 5.5, .35, 18, RGBColor(185, 222, 225), False, margin=0)
textbox(s, "仕事の流れを、前に進める。", .78, 4.05, 6.4, .48, 25, WHITE, True, margin=0)
textbox(s, "タスク・予定・現場の判断をひとつにつなぐ、BrSE向けの仕事の相棒", .8, 4.72, 6.2, .42, 12, RGBColor(177, 201, 211), margin=0)
shape(s, MSO_SHAPE.RECTANGLE, .8, 5.65, 1.3, .04, MINT)
textbox(s, "提案者：System EXE Vietnam", .8, 5.9, 4.2, .28, 11, RGBColor(177, 201, 211), margin=0)
footer(s, 1, True); add_note(s, "最初に、これはAIチャットのデモではなく、BrSEの日々の仕事を前に進めるワークスペースだと伝える。")

# 2. Problem
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s, "01 / WHY", "BrSEの仕事は、情報が分断されている。", "判断するための材料はある。でも、次の一歩に変換する時間が足りない。")
card(s, .75, 2.45, 2.75, 2.5, "タスク", "WBS・個人メモ・チャットに散らばり、今日やることが見えにくい。", CORAL, "01")
card(s, 3.7, 2.45, 2.75, 2.5, "予定", "会議・締切・Google Calendarと、実作業の優先順位が別々に存在する。", GOLD, "02")
card(s, 6.65, 2.45, 2.75, 2.5, "コミュニケーション", "顧客質問やオフショア連携に、翻訳・整理・確認の往復が生まれる。", VIOLET, "03")
card(s, 9.6, 2.45, 2.75, 2.5, "報告", "日報やSOSは重要。でも、書くこと自体が現場の負担になる。", MINT, "04")
textbox(s, "課題は「AIがない」ことではなく、AIが仕事の流れの中にいないこと。", 1.15, 5.65, 11.0, .5, 20, NAVY, True, "Aptos Display", PP_ALIGN.CENTER, margin=0)
footer(s, 2); add_note(s, "課題を機能一覧から始めず、現場の分断から始める。")

# 3. Concept
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, NAVY); title(s, "02 / CONCEPT", "AIは、仕事を奪うのではなく。", "迷いを減らし、人が判断すべきところに集中できる状態をつくる。", True)
textbox(s, "次の一歩を\n見つける。", .78, 2.55, 4.0, 1.28, 36, WHITE, True, "Aptos Display", margin=0)
shape(s, MSO_SHAPE.RECTANGLE, 5.25, 2.62, .05, 2.15, CYAN)
card(s, 5.75, 2.35, 2.15, 2.3, "見える", "タスク・予定・リスクを、同じ画面で把握する。", CYAN, "✓", True)
card(s, 8.1, 2.35, 2.15, 2.3, "整える", "AIに優先順位・文章・差分を整理させる。", MINT, "✓", True)
card(s, 10.45, 2.35, 2.15, 2.3, "動く", "判断したら、そのまま実行・報告につなげる。", GOLD, "✓", True)
textbox(s, "HUMAN DECISION  ×  AI MOMENTUM", 5.78, 5.38, 6.3, .3, 12, RGBColor(170, 201, 211), True, "Aptos", margin=0)
footer(s, 3, True); add_note(s, "3つの原則を、画面構成と機能の設計基準にした。")

# 4. Flow
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s, "03 / EXPERIENCE", "1日の仕事を、ひとつの流れにする。", "入力する → 判断する → 伝える、の往復を短くする設計。")
steps = [("01", "CAPTURE", "WBSを取り込む\nタスクを登録する", CYAN), ("02", "FOCUS", "今日の優先順位を\nAIで整える", MINT), ("03", "EXECUTE", "予定に落とし込み\n進捗を更新する", VIOLET), ("04", "RESPOND", "SOS・顧客質問・\nオフショア差分に対応", CORAL), ("05", "REPORT", "日報を生成し\nチームへ共有する", GOLD)]
for i, (num, head, body, ac) in enumerate(steps):
    x = .78 + i * 2.48
    shape(s, MSO_SHAPE.OVAL, x+.56, 2.55, .76, .76, ac)
    textbox(s, num, x+.56, 2.78, .76, .2, 11, NAVY, True, "Aptos", PP_ALIGN.CENTER, margin=0)
    if i < 4: shape(s, MSO_SHAPE.CHEVRON, x+1.55, 2.82, .55, .2, LINE)
    textbox(s, head, x, 3.65, 1.9, .25, 11, INK, True, "Aptos", PP_ALIGN.CENTER, margin=0)
    textbox(s, body, x, 4.12, 1.9, .55, 11, MUTED, False, "Aptos", PP_ALIGN.CENTER, margin=0)
textbox(s, "毎日の仕事が、最後に「学び」と「共有」に変わる。", 1.2, 5.9, 10.8, .38, 18, NAVY, True, "Aptos Display", PP_ALIGN.CENTER, margin=0)
footer(s, 4); add_note(s, "デモではこの順番に沿って、機能を見せると理解されやすい。")

# 5. Features
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s, "04 / PRODUCT", "ひとつの画面に、6つの仕事の武器。", "AI機能を単独で見せず、現場のアクションに接続している。")
features = [("WBS Import", "CSV / Excel / PDFから、プロジェクトの骨格を取り込む。", CYAN), ("AI Schedule", "期限・優先度・負荷を見て、今日の動きを提案する。", MINT), ("Calendar", "タスクを時間枠へドラッグし、予定に変える。", VIOLET), ("Daily Report", "作業ログから、日報のたたき台を生成する。", GOLD), ("SOS", "停滞やリスクを、要約・メール案まで一気に整理する。", CORAL), ("Offshore Assist", "仕様差分・顧客質問・テスト作成を支援する。", NAVY2)]
for i, (head, body, ac) in enumerate(features):
    x = .78 + (i % 3) * 4.15; y = 2.3 + (i // 3) * 1.85
    card(s, x, y, 3.72, 1.45, head, body, ac, "→")
footer(s, 5); add_note(s, "機能を6つのボタンとしてではなく、仕事の詰まりを解消する6つの武器として説明する。")

# 6. Demo scenario
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, PAPER); title(s, "05 / DEMO STORY", "保険料計算ロジックが、180分止まったら。", "現場で起きる具体的な1シーンから、GROWTH PARTNERの価値を見せる。")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, .8, 2.3, 3.0, 3.2, WHITE, LINE)
textbox(s, "現場の状態", 1.05, 2.62, 2.2, .25, 11, CORAL, True, margin=0)
textbox(s, "PremiumCalculator.java\n\n・タスクが停滞\n・期限への影響が不明\n・PMへの説明材料がない", 1.05, 3.12, 2.4, 1.75, 15, INK, True, margin=0)
shape(s, MSO_SHAPE.CHEVRON, 4.3, 3.55, .75, .62, CYAN)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.45, 2.3, 3.0, 3.2, NAVY, None)
textbox(s, "GROWTH PARTNER", 5.78, 2.62, 2.35, .25, 11, CYAN, True, margin=0)
textbox(s, "SOSを押す\n\nAIが状況を要約\nリスクを分類\n連絡文を下書き", 5.78, 3.12, 2.3, 1.72, 15, WHITE, True, margin=0)
shape(s, MSO_SHAPE.CHEVRON, 8.95, 3.55, .75, .62, MINT)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 10.1, 2.3, 2.45, 3.2, WHITE, LINE)
textbox(s, "次のアクション", 10.38, 2.62, 1.9, .25, 11, MINT, True, margin=0)
textbox(s, "PMへ即時共有\n\n判断に必要な情報が\n1つの画面に揃う", 10.38, 3.12, 1.85, 1.5, 15, INK, True, margin=0)
textbox(s, "AIの回答ではなく、現場の“次の一手”が成果。", 1.0, 6.08, 11.1, .35, 18, NAVY, True, "Aptos Display", PP_ALIGN.CENTER, margin=0)
footer(s, 6); add_note(s, "デモの山場。SOSを押す前後で、現場の不確実性がどう減るかを見せる。")

# 7. UX
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, NAVY); title(s, "06 / UX", "AIっぽさではなく、仕事っぽさをデザインする。", "情報量は多い。でも、目線の順番は迷わせない。", True)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, .85, 2.25, 7.25, 3.75, RGBColor(239, 246, 248), None)
shape(s, MSO_SHAPE.RECTANGLE, .85, 2.25, 1.1, 3.75, NAVY2)
shape(s, MSO_SHAPE.RECTANGLE, 1.95, 2.25, 4.9, 3.75, WHITE)
shape(s, MSO_SHAPE.RECTANGLE, 6.85, 2.25, 1.25, 3.75, RGBColor(243, 247, 248))
textbox(s, "NAV", 1.12, 3.95, .55, .2, 9, CYAN, True, PP_ALIGN.CENTER, margin=0)
textbox(s, "TODAY\n\nTASKS\n\nTOOLS", 1.08, 2.75, .62, 1.25, 9, WHITE, True, PP_ALIGN.CENTER, margin=0)
textbox(s, "EXECUTION DESK", 2.3, 2.58, 3.2, .2, 8, MUTED, True, margin=0)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.3, 3.05, 4.1, .65, NAVY, None)
textbox(s, "今日やることを、前に進める。", 2.55, 3.25, 3.6, .2, 12, WHITE, True, margin=0)
for j, ac in enumerate([CORAL, GOLD, MINT]):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.3, 4.05+j*.46, 4.1, .3, PAPER, LINE)
    shape(s, MSO_SHAPE.RECTANGLE, 2.3, 4.05+j*.46, .05, .3, ac)
textbox(s, "CONTEXT", 7.2, 3.95, .75, .2, 8, MUTED, True, margin=0)
textbox(s, "Detail\nPanel", 7.15, 4.28, .75, .5, 12, INK, True, PP_ALIGN.CENTER, margin=0)
card(s, 8.7, 2.52, 3.55, 1.05, "Hierarchy", "Việc cần làm nằm giữa. Bối cảnh nằm hai bên.", CYAN, "01", True)
card(s, 8.7, 3.82, 3.55, 1.05, "Feedback", "Mỗi thao tác đều có trạng thái rõ: đang làm, trễ, xong.", MINT, "02", True)
card(s, 8.7, 5.12, 3.55, 1.05, "Human tone", "AI là cộng sự trong công việc, không phải nhân vật trình diễn.", GOLD, "03", True)
footer(s, 7, True); add_note(s, "Đây là slide để nói về UX: sản phẩm giảm cognitive load, không chỉ thêm công nghệ.")

# 8. Value
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s, "07 / VALUE", "Từ “đang bận” sang “biết bước tiếp theo”.", "Giá trị lớn nhất là giảm thời gian chuyển đổi giữa các loại công việc.")
textbox(s, "TRƯỚC", 1.05, 2.3, 1.4, .22, 10, CORAL, True, margin=0)
textbox(s, "SAU", 7.15, 2.3, 1.4, .22, 10, MINT, True, margin=0)
for i, t in enumerate(["Mở nhiều công cụ", "Tự gom thông tin", "Tự viết báo cáo", "Phản ứng khi có rủi ro"]):
    y = 2.8 + i*.65; textbox(s, "—  " + t, 1.05, y, 4.5, .28, 15, MUTED, margin=0)
for i, t in enumerate(["Một workspace", "AI gợi ý ưu tiên", "Bản nháp trong vài giây", "SOS có ngữ cảnh"]):
    y = 2.8 + i*.65; textbox(s, "✓  " + t, 7.15, y, 4.5, .28, 15, INK, True, margin=0)
shape(s, MSO_SHAPE.RECTANGLE, 6.3, 2.25, .03, 2.95, LINE)
textbox(s, "Không đo bằng số lượng tính năng.\nĐo bằng số lần người dùng phải tự nối các mảnh thông tin.", 1.05, 5.75, 11.1, .55, 18, NAVY, True, "Aptos Display", PP_ALIGN.CENTER, margin=0)
footer(s, 8); add_note(s, "Không nên hứa KPI chưa đo. Dùng nhóm giá trị định tính và đề xuất đo sau pilot.")

# 9. Tech and trust
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, PAPER); title(s, "08 / BUILD", "Đủ thực tế để demo. Đủ nền tảng để phát triển tiếp.", "Một web app đơn giản ở bề mặt, nhưng có cấu trúc rõ cho giai đoạn pilot.")
layers = [("FRONTEND", "HTML / CSS / JavaScript\nResponsive workspace", CYAN), ("BACKEND", "Spring Boot\nREST API / PDFBox", MINT), ("AI", "Gemini via Spring AI\nUser key per request", VIOLET)]
for i, (head, body, ac) in enumerate(layers):
    x = .95 + i*4.05; card(s, x, 2.45, 3.45, 1.55, head, body, ac, "→")
    if i < 2: shape(s, MSO_SHAPE.CHEVRON, x+3.62, 3.04, .3, .25, LINE)
card(s, .95, 4.55, 3.45, 1.25, "Privacy by design", "API key được lưu ở browser và gửi theo request, không hard-code vào source.", GOLD, "✓")
card(s, 4.98, 4.55, 3.45, 1.25, "Composable", "Các AI operation tách theo nghiệp vụ: schedule, SOS, nippo, offshore.", CYAN, "✓")
card(s, 9.01, 4.55, 3.45, 1.25, "Next pilot", "Đo time-to-report, time-to-triage và mức độ adoption theo team.", MINT, "✓")
footer(s, 9); add_note(s, "Nói ngắn về kỹ thuật để chứng minh tính khả thi, không biến bài trình bày thành code review.")

# 10. Close
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, NAVY)
textbox(s, "GROWTH PARTNER", .78, .8, 5.0, .35, 13, CYAN, True, margin=0)
textbox(s, "AI không làm BrSE\nít quan trọng hơn.\nAI giúp BrSE\nđi xa hơn.", .75, 1.55, 7.6, 2.4, 34, WHITE, True, "Aptos Display", margin=0)
shape(s, MSO_SHAPE.RECTANGLE, .8, 4.5, 1.25, .05, MINT)
textbox(s, "Từ một ngày làm việc nhiều mảnh ghép\nđến một nhịp làm việc có chủ đích.", .8, 4.82, 6.4, .65, 17, RGBColor(190, 215, 220), margin=0)
if MASCOT.exists(): s.shapes.add_picture(str(MASCOT), Inches(9.65), Inches(1.55), width=Inches(2.25), height=Inches(2.25))
textbox(s, "THANK YOU", 9.1, 4.85, 3.4, .38, 22, WHITE, True, "Aptos Display", PP_ALIGN.CENTER, margin=0)
textbox(s, "Demo link: http://localhost:8080", 8.7, 5.45, 4.2, .25, 11, RGBColor(169, 199, 208), False, "Aptos", PP_ALIGN.CENTER, margin=0)
footer(s, 10, True); add_note(s, "Kết thúc bằng một câu ngắn, sau đó chuyển ngay sang demo thực tế.")

prs.save(OUT)
print(OUT.resolve())
